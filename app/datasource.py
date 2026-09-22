"""Menu verisini okuyan katman.

Desteklenen bicimler: .xlsx (Excel), .csv, .json
Kaynak: yerel dosya ya da URL (GitHub raw linki, Google Sheets CSV yayin
linki). URL modunda veri bellekte TTL suresi kadar tutulur; indirme
basarisiz olursa son basarili veri kullanilmaya devam eder (bot susmaz).
"""

from __future__ import annotations

import csv
import io
import json
import logging
import time
from datetime import date, datetime
from typing import Any, Iterable

import httpx

from .config import Config
from .nlu import MONTHS, normalize

log = logging.getLogger(__name__)

# CSV/metin hucrelerinde kabul edilen tarih yazimlari
_DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d")


def parse_cell_date(value: Any) -> date | None:
    """Hucre degerinden tarih cikarir.

    Excel gercek tarih hucresi, "2026-09-01", "01.09.2026" ve Turkce
    yazilmis "1 Eylul 2026" bicimlerinin hepsini anlar.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip()
    if not text:
        return None

    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    # "1 Eylul 2026" / "1 Eylül 2026"
    parts = normalize(text).split()
    if len(parts) == 3 and parts[1] in MONTHS:
        try:
            return date(int(parts[2]), MONTHS[parts[1]], int(parts[0]))
        except ValueError:
            return None
    return None


class MenuStore:
    def __init__(self, cfg: Config):
        self.cfg = cfg
        self._rows: dict[date, dict[str, str]] = {}
        self._loaded_at: float = 0.0
        self._last_error: str | None = None

    # -- genel API ----------------------------------------------------------

    def get(self, day: date) -> dict[str, str] | None:
        self._refresh_if_needed()
        return self._rows.get(day)

    def get_many(self, days: list[date]) -> list[tuple[date, dict[str, str] | None]]:
        self._refresh_if_needed()
        return [(d, self._rows.get(d)) for d in days]

    def reload(self) -> int:
        """Onbellegi zorla tazeler, okunan gun sayisini dondurur."""
        self._load()
        return len(self._rows)

    @property
    def columns(self) -> list[str]:
        """Veride gorulen kolon adlari (config yazarken ise yarar)."""
        for row in self._rows.values():
            return list(row.keys())
        return []

    @property
    def date_span(self) -> tuple[date, date] | None:
        if not self._rows:
            return None
        days = sorted(self._rows)
        return days[0], days[-1]

    @property
    def status(self) -> str:
        if not self._loaded_at:
            return "henüz yüklenmedi"
        age = int(time.time() - self._loaded_at)
        src = self.cfg.data.url if self.cfg.data.source == "url" else str(self.cfg.resolved_local_path())
        line = f"{len(self._rows)} gün | kaynak: {src} | {age} sn önce yüklendi"
        span = self.date_span
        if span:
            line += f"\nKapsam: {span[0]:%d.%m.%Y} – {span[1]:%d.%m.%Y}"
        if self._last_error:
            line += f"\n⚠️ son hata: {self._last_error}"
        return line

    # -- ic isleyis ---------------------------------------------------------

    def _refresh_if_needed(self) -> None:
        ttl = self.cfg.data.cache_ttl_minutes * 60
        if self.cfg.data.source == "local":
            # Yerel dosyada mtime degistiyse yeniden oku.
            path = self.cfg.resolved_local_path()
            mtime = path.stat().st_mtime if path.exists() else 0
            if not self._loaded_at or mtime > self._loaded_at:
                self._load()
            return
        if not self._loaded_at or (time.time() - self._loaded_at) > ttl:
            self._load()

    def _load(self) -> None:
        try:
            raw = self._fetch()
            rows = self._parse(raw)
        except Exception as exc:  # noqa: BLE001 - bot ayakta kalmali
            self._last_error = f"{type(exc).__name__}: {exc}"
            log.warning("Menü yüklenemedi (%s). Eldeki veri kullanılacak.", self._last_error)
            if not self._rows:
                self._loaded_at = time.time()  # surekli denemeyi engelle
            return

        self._rows = rows
        self._loaded_at = time.time()
        self._last_error = None
        log.info("Menü yüklendi: %d gün", len(rows))

    def _fetch(self) -> bytes:
        if self.cfg.data.source == "url":
            url = self.cfg.data.url.strip()
            if not url:
                raise ValueError("config.yaml içinde data.url boş")
            resp = httpx.get(url, timeout=20, follow_redirects=True)
            resp.raise_for_status()
            return resp.content
        path = self.cfg.resolved_local_path()
        if not path.exists():
            raise FileNotFoundError(path)
        return path.read_bytes()

    def _parse(self, raw: bytes) -> dict[date, dict[str, str]]:
        # .xlsx aslinda bir zip; imzasi "PK".
        if raw[:2] == b"PK":
            return self._parse_xlsx(raw)

        text = raw.decode("utf-8-sig", errors="replace")
        stripped = text.lstrip()
        if stripped[:1] in "[{":
            return self._parse_json(stripped)
        return self._parse_csv(text)

    def _parse_xlsx(self, raw: bytes) -> dict[date, dict[str, str]]:
        from openpyxl import load_workbook  # agir import, sadece gerekince

        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        try:
            sheet = self.cfg.data.sheet_name.strip()
            ws = wb[sheet] if sheet else wb.worksheets[0]

            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                return {}
            keys = [("" if h is None else str(h)) for h in header]
            records = [dict(zip(keys, row)) for row in rows]
        finally:
            wb.close()
        return self._rows_from_dicts(records)

    def _parse_csv(self, text: str) -> dict[date, dict[str, str]]:
        try:
            delimiter = csv.Sniffer().sniff(text[:2048], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        return self._rows_from_dicts(reader)

    def _parse_json(self, text: str) -> dict[date, dict[str, str]]:
        payload = json.loads(text)
        if isinstance(payload, dict):
            # {"2026-09-22": {"ogle": "..."} } bicimi
            records = []
            for key, value in payload.items():
                row = dict(value or {})
                row[self.cfg.data.date_column] = key
                records.append(row)
            return self._rows_from_dicts(records)
        return self._rows_from_dicts(payload)

    def _rows_from_dicts(self, records: Iterable[dict]) -> dict[date, dict[str, str]]:
        date_col = self.cfg.data.date_column.strip().lower()
        out: dict[date, dict[str, str]] = {}

        for raw in records:
            if not raw:
                continue
            # Kolon adlarindaki bosluk/buyuk-kucuk farklarini tolere et.
            cells = {(k or "").strip().lower(): v for k, v in raw.items() if k}

            day = parse_cell_date(cells.get(date_col))
            if day is None:
                continue

            out[day] = {
                k: ("" if v is None else str(v).strip()) for k, v in cells.items()
            }
        return out
